from dotenv import load_dotenv
load_dotenv()

import os
import shutil
import pandas as pd

from PIL import Image
import pytesseract

from docx import Document as DocxDocument
from pptx import Presentation

from langchain.document_loaders import PyPDFLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.embeddings.openai import OpenAIEmbeddings
from langchain.vectorstores import Chroma
from langchain.llms import OpenAI
from langchain.chains import RetrievalQA
from langchain.schema import Document


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
VECTOR_DIR = os.path.join(BASE_DIR, "vectorstore")


# -------------------------------
# Tesseract OCR configuration
# -------------------------------
TESSERACT_CMD = os.getenv("TESSERACT_CMD")

if TESSERACT_CMD and os.path.exists(TESSERACT_CMD):
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
else:
    detected_tesseract = shutil.which("tesseract")

    if detected_tesseract:
        pytesseract.pytesseract.tesseract_cmd = detected_tesseract
    else:
        default_tesseract_path = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

        if os.path.exists(default_tesseract_path):
            pytesseract.pytesseract.tesseract_cmd = default_tesseract_path


# -------------------------------
# OpenAI Embedding and LLM
# -------------------------------
embedding = OpenAIEmbeddings(
    model="text-embedding-ada-002"
)

llm = OpenAI(
    model_name="gpt-3.5-turbo-instruct",
    temperature=0,
    max_tokens=500
)


# -------------------------------
# Image OCR
# -------------------------------
def extract_text_from_image(file_path):
    try:
        image = Image.open(file_path)
        text = pytesseract.image_to_string(image, lang="eng")

        if not text.strip():
            return "No readable text was detected in this image."

        return text

    except pytesseract.TesseractNotFoundError:
        raise RuntimeError(
            "Tesseract OCR executable was not found. "
            "Install Tesseract OCR and add it to PATH, or set TESSERACT_CMD in .env. "
            "Example: TESSERACT_CMD=C:\\Program Files\\Tesseract-OCR\\tesseract.exe"
        )


# -------------------------------
# Excel Loader
# -------------------------------
def load_excel_as_documents(file_path):
    docs = []

    excel_data = pd.read_excel(
        file_path,
        sheet_name=None,
        engine="openpyxl"
    )

    for sheet_name, df in excel_data.items():
        df = df.fillna("")

        for index, row in df.iterrows():
            row_text = "\n".join(
                [f"{col}: {row[col]}" for col in df.columns]
            )

            docs.append(
                Document(
                    page_content=f"Sheet: {sheet_name}\nRow: {index + 1}\n{row_text}",
                    metadata={
                        "source": file_path,
                        "sheet_name": sheet_name,
                        "row": index + 1
                    }
                )
            )

    return docs


# -------------------------------
# Word Loader
# -------------------------------
def load_word_document(file_path):
    doc = DocxDocument(file_path)

    text = "\n".join(
        [para.text for para in doc.paragraphs if para.text.strip()]
    )

    if not text.strip():
        text = "No readable text was found in this Word document."

    return [
        Document(
            page_content=text,
            metadata={"source": file_path}
        )
    ]


# -------------------------------
# PowerPoint Loader
# -------------------------------
def load_powerpoint(file_path):
    prs = Presentation(file_path)
    docs = []

    for i, slide in enumerate(prs.slides):
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_parts.append(shape.text)

        slide_text = "\n".join(text_parts)

        if not slide_text.strip():
            slide_text = "No readable text was found on this slide."

        docs.append(
            Document(
                page_content=slide_text,
                metadata={
                    "source": file_path,
                    "slide": i + 1
                }
            )
        )

    return docs


# -------------------------------
# Main Document Router
# -------------------------------
def load_document(file_path):
    ext = os.path.splitext(file_path)[1].lower().replace(".", "")

    if ext == "pdf":
        return PyPDFLoader(file_path).load()

    elif ext == "docx":
        return load_word_document(file_path)

    elif ext == "pptx":
        return load_powerpoint(file_path)

    elif ext == "xlsx":
        return load_excel_as_documents(file_path)

    elif ext in ["png", "jpg", "jpeg"]:
        text = extract_text_from_image(file_path)

        return [
            Document(
                page_content=text,
                metadata={"source": file_path}
            )
        ]

    else:
        raise ValueError(f"Unsupported file type: {ext}")


# -------------------------------
# Process File
# -------------------------------
def process_file(file_path):
    docs = load_document(file_path)

    if docs is None:
        raise ValueError(
            f"No documents were returned from load_document() for file: {file_path}"
        )

    if not isinstance(docs, list):
        raise TypeError(
            f"load_document() must return a list of Document objects. Got: {type(docs)}"
        )

    if len(docs) == 0:
        raise ValueError(
            f"No readable content found in file: {file_path}"
        )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=2000,
        chunk_overlap=200
    )

    chunks = splitter.split_documents(docs)

    if len(chunks) == 0:
        raise ValueError(
            f"Document was loaded but no chunks were created for file: {file_path}"
        )

    vectordb = Chroma.from_documents(
        documents=chunks,
        embedding=embedding,
        persist_directory=VECTOR_DIR
    )

    vectordb.persist()


# -------------------------------
# Ask Question
# -------------------------------
def ask_question(question):
    vectordb = Chroma(
        persist_directory=VECTOR_DIR,
        embedding_function=embedding
    )

    retriever = vectordb.as_retriever()

    qa = RetrievalQA.from_chain_type(
        llm=llm,
        retriever=retriever
    )

    return qa.run(question)